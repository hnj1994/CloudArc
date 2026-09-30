"""CSV / XLSX export of any table and the standard reports (FR-903, FR-904)."""
from __future__ import annotations

import csv
import io
import json
from datetime import date

from .. import budgets as budgets_mod
from ..analytics import costs
from ..analytics.filters import Scope
from ..db import Database
from ..tenants import get_tenant

STANDARD_REPORTS = ["cost_allocation", "resource_explorer", "unit_economics", "budget_vs_actual", "tag_compliance"]


def _flat(v):
    if isinstance(v, (dict, list)):
        return json.dumps(v, default=str)
    if isinstance(v, (date,)):
        return v.isoformat()
    return v


def columns_of(rows: list[dict]) -> list[str]:
    cols: list[str] = []
    for r in rows:
        for k in r:
            if k not in cols:
                cols.append(k)
    return cols


def to_csv(rows: list[dict]) -> bytes:
    buf = io.StringIO()
    cols = columns_of(rows)
    w = csv.writer(buf)
    w.writerow(cols)
    for r in rows:
        w.writerow([_flat(r.get(c)) for c in cols])
    return buf.getvalue().encode("utf-8-sig")  # BOM so Excel opens ₹ and UTF-8 correctly


def to_xlsx(rows: list[dict], title: str = "Export") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    cols = columns_of(rows)
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DCE6F5")
    for r in rows:
        ws.append([_flat(r.get(c)) for c in cols])
    for i, c in enumerate(cols, start=1):
        width = max([len(str(c))] + [len(str(_flat(r.get(c)) or "")) for r in rows[:200]])
        ws.column_dimensions[ws.cell(1, i).column_letter].width = min(max(10, width + 2), 60)
        if any(k in c for k in ("cost", "amount", "actual", "forecast", "saving")):
            for cell in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                cell[0].number_format = "#,##,##0.00"
    ws.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def standard_report(db: Database, name: str, scope: Scope) -> list[dict]:
    if name == "cost_allocation":
        return costs.group_by(db, scope, ["account", "resource_group", "service", "location"])
    if name == "resource_explorer":
        return costs.resource_explorer(db, scope, limit=100000)
    if name == "unit_economics":
        return costs.unit_economics(db, scope)
    if name == "budget_vs_actual":
        t = get_tenant(db, scope.tenant_id)
        rows = []
        for b in budgets_mod.list_budgets(db, scope.tenant_id):
            s = budgets_mod.status(db, scope.tenant_id, b, scope.date_to, t["fiscal_year_start"])
            rows.append({k: s[k] for k in ("name", "scope_type", "scope_value", "period", "period_start", "period_end", "as_of",
                                           "is_partial", "amount", "actual", "actual_pct", "forecast", "forecast_pct")}
                        | {"calibration": s["calibration"]["status"]})
        return rows
    if name == "tag_compliance":
        t = get_tenant(db, scope.tenant_id)
        cov = costs.tag_coverage(db, scope, t["required_tags"], top_untagged=100000)
        return [{"type": "coverage", **c} for c in cov["required_tags"]] + [{"type": "violation", **v} for v in cov["violations"]]
    raise ValueError(f"unknown report; expected one of {STANDARD_REPORTS}")
