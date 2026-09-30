"""DOCX and PDF renderers for :class:`~cloudarc.reports.builder.Report`."""
from __future__ import annotations

import os
from pathlib import Path

from ..config import get_settings
from .builder import Report

ACCENT = (0x1F, 0x4E, 0x9E)


# ---- DOCX ---------------------------------------------------------------------------------------

def _docx_page_field(paragraph, instr: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, instr), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def to_docx(report: Report, path: Path) -> Path:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    for s in ("Heading 1", "Heading 2"):
        doc.styles[s].font.color.rgb = RGBColor(*ACCENT)
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(0.8)
    header = sec.header.paragraphs[0]
    header.text = f"{report.title} — {report.tenant}"
    header.runs[0].font.size = Pt(8)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run(f"{get_settings().org_name}   |   Page ")
    _docx_page_field(footer, "PAGE")
    footer.add_run(" of ")
    _docx_page_field(footer, "NUMPAGES")
    for r in footer.runs:
        r.font.size = Pt(8)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = t.add_run(report.title)
    run.bold, run.font.size, run.font.color.rgb = True, Pt(22), RGBColor(*ACCENT)
    st = doc.add_paragraph().add_run(report.subtitle)
    st.font.size = Pt(13)

    def table(columns, rows, numeric=()):
        tbl = doc.add_table(rows=1, cols=len(columns))
        tbl.style = "Light Grid Accent 1"
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, c in enumerate(columns):
            cell = tbl.rows[0].cells[i]
            cell.text = str(c)
            cell.paragraphs[0].runs[0].bold = True
        for row in rows:
            cells = tbl.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = "" if v is None else str(v)
                if i in numeric:
                    cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(8.5)
        doc.add_paragraph()

    for block in report.blocks:
        kind = block[0]
        if kind == "h1":
            doc.add_heading(block[1], level=1)
        elif kind == "h2":
            doc.add_heading(block[1], level=2)
        elif kind == "p":
            doc.add_paragraph(block[1])
        elif kind == "bullets":
            for item in block[1]:
                doc.add_paragraph(item, style="List Bullet")
        elif kind == "numbered":
            for item in block[1]:
                doc.add_paragraph(item, style="List Number")
        elif kind == "kv":
            table(["Item", "Value"], [[k, v] for k, v in block[1]])
        elif kind == "table":
            spec = block[1]
            if spec["rows"]:
                table(spec["columns"], spec["rows"], spec["numeric"])
        elif kind == "image":
            doc.add_picture(block[1], width=Inches(block[2]))
        elif kind == "pagebreak":
            doc.add_page_break()
    doc.save(path)
    return path


# ---- PDF ----------------------------------------------------------------------------------------

def _register_fonts() -> tuple[str, str]:
    import matplotlib
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    ttf = os.path.join(matplotlib.get_data_path(), "fonts", "ttf")
    if "DejaVuSans" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("DejaVuSans", os.path.join(ttf, "DejaVuSans.ttf")))  # has the ₹ glyph
        pdfmetrics.registerFont(TTFont("DejaVuSans-Bold", os.path.join(ttf, "DejaVuSans-Bold.ttf")))
    return "DejaVuSans", "DejaVuSans-Bold"


def to_pdf(report: Report, path: Path) -> Path:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import (Image, ListFlowable, ListItem, PageBreak, Paragraph, SimpleDocTemplate, Spacer,
                                    Table, TableStyle)
    from xml.sax.saxutils import escape

    regular, bold = _register_fonts()
    ss = getSampleStyleSheet()
    accent = colors.Color(*(c / 255 for c in ACCENT))
    body = ParagraphStyle("body", parent=ss["BodyText"], fontName=regular, fontSize=9.5, leading=13)
    small = ParagraphStyle("small", parent=body, fontSize=8, leading=10)
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontName=bold, fontSize=15, textColor=accent, spaceBefore=10)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontName=bold, fontSize=11.5, textColor=accent, spaceBefore=6)
    title = ParagraphStyle("title", parent=h1, fontSize=22, leading=28, spaceAfter=6)

    def tbl(columns, rows, numeric=()):
        data = [[Paragraph(f"<b>{escape(str(c))}</b>", small) for c in columns]]
        for row in rows:
            data.append([Paragraph(escape("" if v is None else str(v)), ParagraphStyle("c", parent=small, alignment=2 if i in numeric else 0))
                         for i, v in enumerate(row)])
        t = Table(data, repeatRows=1, hAlign="LEFT")
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.Color(0.9, 0.93, 0.98)),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.Color(0.75, 0.78, 0.85)),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        return [t, Spacer(1, 8)]

    story = [Spacer(1, 1.2 * inch), Paragraph(escape(report.title), title), Paragraph(escape(report.subtitle), h2), Spacer(1, 12)]
    for block in report.blocks:
        kind = block[0]
        if kind == "h1":
            story.append(Paragraph(escape(block[1]), h1))
        elif kind == "h2":
            story.append(Paragraph(escape(block[1]), h2))
        elif kind == "p":
            story.append(Paragraph(escape(block[1]), body))
        elif kind in ("bullets", "numbered"):
            story.append(ListFlowable([ListItem(Paragraph(escape(i), body)) for i in block[1]],
                                      bulletType="bullet" if kind == "bullets" else "1", bulletFontName=regular, leftIndent=14))
        elif kind == "kv":
            story += tbl(["Item", "Value"], [[k, v] for k, v in block[1]])
        elif kind == "table" and block[1]["rows"]:
            story += tbl(block[1]["columns"], block[1]["rows"], block[1]["numeric"])
        elif kind == "image":
            img = Image(block[1])
            ratio = img.imageHeight / img.imageWidth
            img.drawWidth, img.drawHeight = block[2] * inch, block[2] * inch * ratio
            story.append(img)
        elif kind == "pagebreak":
            story.append(PageBreak())

    def decorate(canvas, doc):
        canvas.saveState()
        canvas.setFont(regular, 7.5)
        canvas.setFillColor(colors.grey)
        canvas.drawString(doc.leftMargin, A4[1] - 0.45 * inch, f"{report.title} — {report.tenant}")
        canvas.drawRightString(A4[0] - doc.rightMargin, 0.45 * inch, f"{get_settings().org_name}   |   Page {doc.page}")
        canvas.restoreState()

    SimpleDocTemplate(str(path), pagesize=A4, leftMargin=0.75 * inch, rightMargin=0.75 * inch,
                      topMargin=0.75 * inch, bottomMargin=0.75 * inch,
                      title=report.title, author="CloudArc").build(story, onFirstPage=decorate, onLaterPages=decorate)
    return path
