from datetime import date, timedelta

import pytest

from cloudarc.analytics import costs
from cloudarc.analytics.filters import Scope
from cloudarc.recommendations import engine
from cloudarc.recommendations.engine import LifecycleError
from cloudarc.recommendations.rules import parse_vm_size, smaller_size
from cloudarc.reports import boq, builder, render
from tests.conftest import AS_OF


def recs_by_category(db, tenant):
    out = {}
    for r in engine.list_recs(db, tenant):
        out.setdefault(r["category"], []).append(r)
    return out


def test_vm_size_parsing():
    assert parse_vm_size("Standard_B4als_v2") == {"family": "B", "vcpus": 4, "features": "als", "version": "v2"}
    assert parse_vm_size("B4als v2")["vcpus"] == 4
    assert parse_vm_size("Standard_E16as_v5")["family"] == "E"
    assert smaller_size(parse_vm_size("Standard_D48s_v5"))["vcpus"] == 20


def test_expected_recommendations_with_evidence(seeded):
    db, _ = seeded
    a = recs_by_category(db, "demo-a")
    vm = a["vm_rightsizing"][0]
    assert "B2als_v2" in vm["title"]
    assert vm["evidence"]["metric_days"] >= 14 and vm["evidence"]["cpu_avg_pct"] < 15
    assert vm["est_monthly_saving"] == pytest.approx(7303.08 / 2, rel=0.05)
    assert a["sql_tier"][0]["evidence"]["target_tier"] == "Basic"
    assert a["orphaned_disk"][0]["resource_name"] == "app-prod-vm01_datadisk_old"
    b = recs_by_category(db, "demo-b")
    assert b["reserved_instance"][0]["resource_name"] == "hp-db-vm01"  # on-demand, 24x7
    assert all(r["resource_name"] != "hp-app-vm01" for r in b["reserved_instance"])  # already reserved
    assert {r["resource_name"] for r in b["disk_tier"]} == {"hp-app-vm01_data01", "hp-db-vm01_data02", "hp-app-vm01_OsDisk"}
    assert "hp-db-vm01_data01" not in {r["resource_name"] for r in b["disk_tier"]}  # busy disk stays premium
    for r in engine.list_recs(db, "demo-b"):
        assert r["evidence"] and r["confidence"] in {"high", "medium", "low"}


def test_no_metrics_no_rightsizing(seeded):
    db, _ = seeded
    db.execute("DELETE FROM resource_metrics WHERE tenant_id = 'demo-a'")
    engine.run(db, "demo-a", AS_OF)
    a = recs_by_category(db, "demo-a")
    assert all(r["status"] == "resolved" for r in a["vm_rightsizing"] + a["sql_tier"])
    assert a["orphaned_disk"][0]["status"] == "open"  # inventory-based rules do not need metrics


def test_lifecycle_and_dismissed_items_stay_dismissed(seeded):
    db, _ = seeded
    rec = recs_by_category(db, "demo-a")["orphaned_disk"][0]
    with pytest.raises(LifecycleError):
        engine.transition(db, "demo-a", rec["id"], "verified")
    with pytest.raises(LifecycleError):
        engine.transition(db, "demo-a", rec["id"], "dismissed")  # reason required
    engine.transition(db, "demo-a", rec["id"], "dismissed", "kept for audit")
    engine.run(db, "demo-a", AS_OF)
    assert db.scalar("SELECT status FROM recommendations WHERE id = ?", [rec["id"]]) == "dismissed"
    with pytest.raises(LifecycleError):
        engine.transition(db, "demo-b", rec["id"], "open")  # other tenant cannot touch it


def test_realized_savings_verification(seeded):
    db, _ = seeded
    rec = recs_by_category(db, "demo-a")["orphaned_disk"][0]
    engine.transition(db, "demo-a", rec["id"], "accepted")
    engine.transition(db, "demo-a", rec["id"], "implemented")
    impl = AS_OF - timedelta(days=20)
    db.execute("UPDATE recommendations SET implemented_at = ? WHERE id = ?", [impl, rec["id"]])
    db.execute("DELETE FROM cost_records WHERE tenant_id = 'demo-a' AND resource_id = ? AND charge_date > ?", [rec["resource_id"], impl])
    assert engine.verify_realized_savings(db, "demo-a", AS_OF) == 1
    row = db.one("SELECT status, realized_monthly_saving FROM recommendations WHERE id = ?", [rec["id"]])
    assert row["status"] == "verified" and row["realized_monthly_saving"] == pytest.approx(5.02 * 30.4, rel=0.01)


def test_advisor_items_merge_into_native(seeded):
    db, _ = seeded
    vm = recs_by_category(db, "demo-a")["vm_rightsizing"][0]
    other = "/subscriptions/x/resourcegroups/rg/providers/microsoft.compute/virtualmachines/other"
    res = engine.merge_advisor(db, "demo-a", None, [
        {"properties": {"resourceMetadata": {"resourceId": vm["resource_id"]}, "shortDescription": {"problem": "Right-size", "solution": "Resize"},
                        "extendedProperties": {"annualSavingsAmount": "12000"}}},
        {"properties": {"resourceMetadata": {"resourceId": other}, "shortDescription": {"problem": "Shut down idle VM", "solution": "Stop it"},
                        "extendedProperties": {"annualSavingsAmount": "2400"}, "recommendationTypeId": "abc"}},
    ])
    assert res == {"merged_into_native": 1, "inserted": 1}
    merged = next(r for r in engine.list_recs(db, "demo-a") if r["id"] == vm["id"])
    assert merged["evidence"]["azure_advisor"]["est_monthly_saving"] == 1000


def test_report_reconciles_with_dashboard_and_renders(seeded, tmp_path):
    db, _ = seeded
    rep = builder.build(db, "demo-b", workdir=tmp_path)
    assert rep.month == "2026-08"
    assert rep.figures["month_total"] == costs.month_view(db, Scope("demo-b"), date(2026, 8, 1))["total"]
    s = costs.summary(db, Scope("demo-b"))
    assert rep.figures["current_mtd"] == s["current"]["amount"] and rep.figures["forecast"] == s["forecast"]["expected"]
    headings = [b[1] for b in rep.blocks if b[0] == "h1"]
    assert [h.split(".")[0] for h in headings] == [str(i) for i in range(1, 11)]
    docx = render.to_docx(rep, tmp_path / "r.docx")
    pdf = render.to_pdf(rep, tmp_path / "r.pdf")
    assert docx.stat().st_size > 20000 and pdf.read_bytes()[:4] == b"%PDF"
    from docx import Document

    text = "\n".join(p.text for p in Document(str(docx)).paragraphs)
    assert "Executive Summary" in text and "Optimization Recommendations" in text


def test_boq_parser_reads_pricing_calculator_export(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Microsoft Azure Estimate"])
    ws.append(["Service category", "Service type", "Custom name", "Region", "Description", "Estimated monthly cost", "Estimated upfront cost"])
    ws.append(["Compute", "Virtual Machines", "App VM", "Central India", "E8s v5", 20373.92, 0])
    ws.append(["Networking", "VPN Gateway", "", "Central India", "VpnGw1AZ", 13943.40, 0])
    ws.append(["Support", "Support", None, None, None, 0, 0])
    ws.append(["Total", None, None, None, None, 34317.32, 0])
    ws.append(["Compute", "ignored after total", None, None, None, 999, 0])
    wb.save(tmp_path / "boq.xlsx")
    rows = boq.parse_calculator_xlsx(tmp_path / "boq.xlsx")
    assert [r["service_type"] for r in rows] == ["Virtual Machines", "VPN Gateway", "Support"]
    assert sum(r["monthly_cost"] for r in rows) == pytest.approx(34317.32)


def test_boq_variance_covers_all_actual_spend(seeded):
    db, _ = seeded
    v = boq.variance(db, "demo-b", "Approved BOQ (March 2026)", date(2026, 8, 1))
    assert v["actual_total"] == pytest.approx(costs.total(db, Scope("demo-b", date(2026, 8, 1), date(2026, 8, 31))), abs=0.05)
    not_in_boq = {line["component"] for line in v["lines"] if not line["in_boq"]}
    assert "Storage – Standard Page Blob v2" in not_in_boq
    disks = next(line for line in v["lines"] if line["component"] == "Managed Disks")
    assert set(disks["actual_by_region"]) == {"Central India", "South India"}
